"""Import the completed TS-RAG/Chronos-Bolt workbook without inference."""

from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path
from typing import Optional

from evaluation.io import write_json, write_variable_metrics
from evaluation.metrics import benchmark_summaries, dataset_summaries, pairwise_comparison
from evaluation.schema import VariableMetric


DATASET_NAMES = {
    "ETTh1": "ETTh1", "ETTh2": "ETTh2", "ETTm1": "ETTm1", "ETTm2": "ETTm2",
    "Weather": "weather", "Electricity": "electricity", "Exchange": "exchange_rate",
}
METHOD_NAMES = {"TS-RAG": "tsrag", "Chronos-Bolt": "chronos_bolt"}
SUMMARY_DATASET_NAMES = {
    "ETTh1": "ETTh1", "ETTh2": "ETTh2", "ETTm1": "ETTm1", "ETTm2": "ETTm2",
    "Weather": "weather", "Electricity": "electricity", "Exchange": "exchange_rate",
}
FREQUENCIES = {
    "ETTh1": "1 hour", "ETTh2": "1 hour", "ETTm1": "15 minutes",
    "ETTm2": "15 minutes", "weather": "10 minutes",
    "electricity": "1 hour", "exchange_rate": "1 day",
}
VARIABLE_NAMES = {
    "SWDR (W/m�)": "SWDR (W/m²)",
    "PAR (�mol/m�/s)": "PAR (μmol/m²/s)",
    "max. PAR (�mol/m�/s)": "max. PAR (μmol/m²/s)",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def import_workbook(source: Path, output_dir: Path, dlinear_source: Optional[Path] = None) -> dict[str, object]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Install openpyxl to import the source workbook") from exc

    workbook = load_workbook(source, data_only=True, read_only=False)
    if "Long data" not in workbook.sheetnames:
        raise ValueError("Workbook is missing the 'Long data' sheet")
    sheet = workbook["Long data"]
    headers = [cell.value for cell in sheet[1]]
    expected = ["Dataset", "Frequency", "Feature ID", "Variable", "Model", "MSE", "MAE", "Test windows", "Horizon"]
    if headers != expected:
        raise ValueError(f"Unexpected Long data columns: {headers}")

    rows: list[VariableMetric] = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        if all(value is None for value in values):
            continue
        dataset, frequency, feature_id, variable, model, mse, mae, windows, horizon = values
        try:
            canonical_dataset = DATASET_NAMES[str(dataset)]
            canonical_method = METHOD_NAMES[str(model)]
        except KeyError as exc:
            raise ValueError(f"Unknown workbook label: {exc.args[0]}") from exc
        rows.append(VariableMetric(
            benchmark_id="benchmark_v1", method=canonical_method,
            dataset=canonical_dataset, frequency=str(frequency), feature_id=int(feature_id),
            variable=VARIABLE_NAMES.get(str(variable), str(variable)), seed=2021, context_length=512,
            prediction_length=int(horizon), test_windows=int(windows),
            mse=float(mse), mae=float(mae), source=source.name,
        ))

    if len(rows) != 756:
        raise ValueError(f"Expected 756 workbook rows, found {len(rows)}")

    if dlinear_source is not None:
        with dlinear_source.open(newline="", encoding="utf-8") as handle:
            reader = csv.DictReader(handle)
            expected_columns = ["dataset", "feature_id", "variable", "mse", "mae", "sample_count", "value_count"]
            if reader.fieldnames != expected_columns:
                raise ValueError(f"Unexpected DLinear columns: {reader.fieldnames}")
            for raw in reader:
                dataset = raw["dataset"]
                rows.append(VariableMetric(
                    benchmark_id="benchmark_v1", method="dlinear", dataset=dataset,
                    frequency=FREQUENCIES[dataset], feature_id=int(raw["feature_id"]),
                    variable=VARIABLE_NAMES.get(raw["variable"], raw["variable"]), seed=2021,
                    context_length=512, prediction_length=64,
                    test_windows=int(raw["sample_count"]), mse=float(raw["mse"]),
                    mae=float(raw["mae"]), source="dlinear/variable_metrics.csv",
                ))
        if sum(row.method == "dlinear" for row in rows) != 378:
            raise ValueError("Expected 378 DLinear variable rows")
    keys = {(row.method, row.dataset, row.feature_id) for row in rows}
    if len(keys) != len(rows):
        raise ValueError("Duplicate method/dataset/feature rows in workbook")

    dataset_rows = dataset_summaries(rows)
    by_dataset_method = {
        (str(row["dataset"]), str(row["method"])): row for row in dataset_rows
    }
    summary_sheet = workbook["Summary"]
    for values in summary_sheet.iter_rows(min_row=7, max_row=13, values_only=True):
        dataset_label, variable_count, total_windows = values[:3]
        dataset = SUMMARY_DATASET_NAMES[str(dataset_label)]
        expected = {
            "tsrag": (float(values[3]), float(values[7])),
            "chronos_bolt": (float(values[4]), float(values[8])),
        }
        for method, (mse, mae) in expected.items():
            actual = by_dataset_method[(dataset, method)]
            if int(actual["variables"]) != int(variable_count):
                raise ValueError(f"Variable count mismatch for {method}/{dataset}")
            if int(actual["test_windows"]) != int(total_windows):
                raise ValueError(f"Test-window count mismatch for {method}/{dataset}")
            if abs(float(actual["mse"]) - mse) > 1e-12 or abs(float(actual["mae"]) - mae) > 1e-12:
                raise ValueError(f"Summary metric mismatch for {method}/{dataset}")
    if dlinear_source is not None:
        native_summary = dlinear_source.parent / "metrics.json"
        native = json.loads(native_summary.read_text(encoding="utf-8"))
        for expected in native["datasets"]:
            actual = by_dataset_method[(expected["dataset"], "dlinear")]
            dataset_native = json.loads(
                (dlinear_source.parent / expected["dataset"] / "metrics.json").read_text(encoding="utf-8")
            )
            if abs(float(actual["mse"]) - float(expected["dlinear_mse"])) > 1e-12:
                raise ValueError(f"DLinear MSE mismatch for {expected['dataset']}")
            if abs(float(actual["mae"]) - float(expected["dlinear_mae"])) > 1e-12:
                raise ValueError(f"DLinear MAE mismatch for {expected['dataset']}")
            if int(actual["variables"]) != int(dataset_native["channel_count"]):
                raise ValueError(f"DLinear variable count mismatch for {expected['dataset']}")
            if int(actual["test_windows"]) != int(dataset_native["sample_count_tsrag_semantics"]):
                raise ValueError(f"DLinear sample count mismatch for {expected['dataset']}")
    benchmark_rows = benchmark_summaries(dataset_rows)
    comparisons = pairwise_comparison(dataset_rows, baseline="chronos_bolt")
    write_variable_metrics(output_dir / "variable_metrics.csv", rows)
    write_json(output_dir / "dataset_metrics.json", dataset_rows)
    write_json(output_dir / "summary.json", {
        "benchmark_id": "benchmark_v1",
        "schema_version": "1.0",
        "method_metadata": {
            "tsrag": {"training_regime": "zero_shot"},
            "chronos_bolt": {"training_regime": "zero_shot"},
            "dlinear": {"training_regime": "supervised_per_dataset"},
        },
        "sources": [
            {"path": source.name, "sha256": sha256(source)},
            *([{"path": "dlinear/variable_metrics.csv", "sha256": sha256(dlinear_source)}] if dlinear_source else []),
        ],
        "row_count": len(rows),
        "benchmark": benchmark_rows,
        "comparisons": comparisons,
    })
    return {"rows": rows, "datasets": dataset_rows, "benchmark": benchmark_rows}
